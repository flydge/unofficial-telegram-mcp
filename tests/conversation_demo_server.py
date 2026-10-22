"""Conversation demo server (0.54 item E): the real MCP server over STDIO against a synthetic account.

Dev and test only; nothing here reaches Telegram. `build_server` serves the complete surface against
`DemoService`, an in-process fake of the broker client that returns the real pydantic response models
for the small account in `conversation_account.py`: Saved Messages, a private chat, a forum
supergroup with topics, a quiet group and an archived channel, with links, mentions and tiny
PDF/XLSX/PPTX/DOCX attachments generated at start-up. Attachments go through a private temporary
artifact cache and the real page, spreadsheet, presentation and document readers. Reply sources are
projected by the real `source_from_message`, so `capability_disabled` names exactly what the real
broker would. Owner approvals are scripted (approve, decline, timeout) after the real approval
renderer has produced the dialog document; a document that cannot be rendered is
`dialog_unavailable`, as on the Mac.

Every tools/call, approval and send attempt is appended to an owner-only JSONL transcript that
`conversation_rubrics.py` scores, so the official-SDK tests and real agents share one rubric.

    python tests/conversation_demo_server.py --transcript ABSOLUTE.jsonl [--policy demo|no_reply_media]
        [--approval approve|decline|timeout] [--send-outcome sent|outcome_unknown]
        [--status-after-unknown sent|outcome_unknown]

Run as a script it first moves HOME into a fresh owner-only temporary directory, before any
`telegram_search_mcp` import, so the live runtime tree is never read or written.
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path


def _private_workdir() -> Path:
    work = Path(tempfile.mkdtemp(prefix="tgmcp-demo-")).resolve()
    os.chmod(work, 0o700)
    (work / "home").mkdir(mode=0o700)
    return work


if __name__ == "__main__":
    # config.py derives runtime paths from HOME at import time.
    _WORKDIR: Path | None = _private_workdir()
    os.environ["HOME"] = str(_WORKDIR / "home")
else:
    _WORKDIR = None

import argparse  # noqa: E402
import atexit  # noqa: E402
import hashlib  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
import secrets  # noqa: E402
import shutil  # noqa: E402
import threading  # noqa: E402
from contextlib import nullcontext  # noqa: E402
from dataclasses import dataclass  # noqa: E402
from datetime import datetime, timedelta, timezone  # noqa: E402

from mcp.server.mcpserver.exceptions import ToolError  # noqa: E402

from telegram_search_mcp.approval_prompt import DETAILS as APPROVAL_DETAILS, NOT_GIVEN, render_approval  # noqa: E402
from telegram_search_mcp.artifact_store import ArtifactStore  # noqa: E402
from telegram_search_mcp.attachment_pages import AttachmentPageReader  # noqa: E402
from telegram_search_mcp.boolean_query import boolean_matches, boolean_seeds, normalize_search_text  # noqa: E402
from telegram_search_mcp.config import RuntimePolicy  # noqa: E402
from telegram_search_mcp.document_reader import read_document  # noqa: E402
from telegram_search_mcp.draft_models import (CancelDraftResponse, DraftPreview, DraftSummary,  # noqa: E402
                                              GetDraftResponse, ListDraftsResponse, ReviseDraftResponse)
from telegram_search_mcp.outgoing_drafts import (_MIME_TYPE, _safe_caption, _safe_display_name,  # noqa: E402
                                                 _safe_message_text)
from telegram_search_mcp.presentations import PresentationReader  # noqa: E402
from telegram_search_mcp.reply_artifact_drafts import (GetReplyArtifactDraftResponse,  # noqa: E402
    PrepareReplyArtifactSendResponse, ReplyArtifactDraftPreview, ReviseReplyArtifactDraftResponse)
from telegram_search_mcp.reply_drafts import (CAPABILITY_DISABLED, PENDING, PREPARED, UNAVAILABLE,  # noqa: E402
    GetReplyDraftResponse, PrepareReplyTextSendResponse, ReplyDraftPreview, ReplyTextDraftPreview,
    ReviseReplyDraftResponse, source_from_message)
from telegram_search_mcp.sanitize import render_evidence, safe_telegram_link  # noqa: E402
from telegram_search_mcp.schemas import (  # noqa: E402
    AttachmentImage, AttachmentResponse, BooleanSearchQuery, CatalogCoverage, CatalogLaneCoverage, ChatListCoverage,
    ChatListResult, ChatListScope, ContextEvidence, Coverage, CreateLocalArtifactResponse, DiscoveryCandidate,
    DiscoveryCoverage, EvidenceAnchor, FileEvidence, ForumTopicMetadata, ForumTopicName, ForumTopicReference,
    ForumTopicResult, GlobalMessageCoverage, GlobalMessageEvidence, GlobalMessageLaneCoverage, HistoryScope,
    ListChatsResponse, ListTopicsResponse, MessageContextResponse, MessageReadResult, MetadataEvidence,
    PrepareArtifactSendResponse, PrepareTextSendResponse, ReadAttachmentResponse, ReadHistoryResponse,
    ReadMessagesResponse, ReadTopicHistoryResponse, SearchBranchCoverage, SearchMatch, SearchMessageResult,
    SearchMessagesResponse, SearchMessagesScope, SearchResponse, SelectedMessage, SelectedMessageAnchor,
    SelectedMessageSender, SelectedMessageText, SelectedMessageTopic, SendPreparedArtifactResponse,
    SendPreparedTextResponse, SourceEvidence, TargetDiscoveryResponse, TopicHistoryScope, TopicScope)
from telegram_search_mcp.send_status_models import send_status  # noqa: E402
from telegram_search_mcp.server import available_tools, build_server  # noqa: E402
from telegram_search_mcp.spreadsheets import SpreadsheetReader  # noqa: E402
from telegram_search_mcp.target_resolver import TargetResolver  # noqa: E402
from telegram_search_mcp.tdjson import MessageNotFound  # noqa: E402

import conversation_account as account  # noqa: E402

UTC = timezone.utc
# Everything the scenarios use, and nothing the fake does not implement: no selections, verified
# targets, multi-chat search, reply chains, styles, topic sends, forwards, video, archive or edits.
DEMO_CAPABILITIES = ("artifacts", "read", "send", "read_messages", "read_history", "list_topics",
                     "read_topic_history", "search_messages", "list_chats", "attachment_pages", "spreadsheets",
                     "presentations", "reply_text_send", "reply_artifact_send", "reply_media_targets")
POLICIES = {
    "demo": DEMO_CAPABILITIES,
    "no_reply_media": tuple(name for name in DEMO_CAPABILITIES if name != "reply_media_targets"),
}
APPROVALS = {"approve": "approved", "decline": "declined", "timeout": "timeout"}
DRAFT_TTL = timedelta(seconds=900)
UNSUPPORTED = "demo_unsupported: the synthetic demo account does not implement this tool"


def demo_policy(name: str) -> RuntimePolicy:
    return RuntimePolicy(enabled_capabilities=POLICIES[name])


# --- synthetic files -------------------------------------------------------------------------------

def _pdf(document: account.Document) -> bytes:
    import fitz
    with fitz.open() as pdf:
        for lines in document.parts:
            page = pdf.new_page(width=420, height=300)
            for index, line in enumerate(lines):
                page.insert_text((24, 40 + 26 * index), line, fontsize=12)
        return pdf.tobytes(no_new_id=True)


def _xlsx(document: account.Document) -> bytes:
    from xml.sax.saxutils import escape
    from xlsx_fixtures import worksheet, xlsx_parts, write_xlsx
    rows = []
    for row_index, row in enumerate(document.parts, 1):
        cells = []
        for column, value in enumerate(row):
            address = "ABCDEFGH"[column] + str(row_index)
            cells.append(f'<c r="{address}"><v>{value}</v></c>' if isinstance(value, int) else
                         f'<c r="{address}" t="inlineStr"><is><t>{escape(value)}</t></is></c>')
        rows.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    buffer = io.BytesIO()
    write_xlsx(buffer, xlsx_parts([("Budget 2026", "visible", worksheet("".join(rows)))]))
    return buffer.getvalue()


def _pptx(document: account.Document) -> bytes:
    import zipfile
    from pptx_fixtures import pptx_parts, text_shape
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_STORED) as archive:
        for name, content in pptx_parts([dict(body=text_shape(title)) for title in document.parts]).items():
            archive.writestr(name, content)
    return buffer.getvalue()


def _docx(document: account.Document) -> bytes:
    from docx import Document
    word = Document()
    for paragraph in document.parts:
        word.add_paragraph(paragraph)
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def synthetic_file(document: account.Document) -> bytes:
    return {"pdf": _pdf, "xlsx": _xlsx, "pptx": _pptx, "docx": _docx}[document.kind](document)


# --- transcript ------------------------------------------------------------------------------------

class Transcript:
    """Append-only owner-only JSONL; one row per tools/call, approval or send attempt."""

    def __init__(self, path: Path | None):
        self._path = path
        self._lock = threading.Lock()
        self._seq = 0
        if path is not None:
            os.close(os.open(path, os.O_WRONLY | os.O_CREAT | os.O_APPEND | os.O_NOFOLLOW, 0o600))

    def write(self, row: dict) -> None:
        with self._lock:
            self._seq += 1
            if self._path is None:
                return
            line = json.dumps({"seq": self._seq, **row}, ensure_ascii=False, sort_keys=True, default=str)
            descriptor = os.open(self._path, os.O_WRONLY | os.O_APPEND | os.O_NOFOLLOW)
            try:
                os.write(descriptor, (line + "\n").encode("utf-8"))
            finally:
                os.close(descriptor)


def record_tool_calls(server, transcript: Transcript) -> None:
    """Wrap the server's one call_tool entry so every call, refusal included, reaches the transcript."""
    original = server.call_tool

    async def call_tool(name, arguments, context=None):
        try:
            result = await original(name, arguments, context)
        except Exception as error:
            transcript.write({"event": "tool_call", "tool": name, "arguments": arguments, "is_error": True,
                              "structured": None, "text": str(error)[:4000], "images": 0})
            raise
        blocks = list(getattr(result, "content", None) or [])
        texts = [block.text for block in blocks if getattr(block, "type", None) == "text"]
        transcript.write({"event": "tool_call", "tool": name, "arguments": arguments,
                          "is_error": bool(getattr(result, "is_error", False)),
                          "structured": getattr(result, "structured_content", None),
                          "text": texts[0][:4000] if texts and getattr(result, "is_error", False) else None,
                          "images": sum(1 for block in blocks if getattr(block, "type", None) == "image")})
        return result

    server.call_tool = call_tool


# --- provider-shaped projections -------------------------------------------------------------------

_ENTITY = re.compile(r"https?://[^\s]+|@[A-Za-z0-9_]{5,32}")


def _entities(text: str) -> list[dict]:
    """TDLib-shaped mention and URL entities (ASCII texts, so UTF-16 offsets equal str offsets)."""
    return [{"@type": "textEntity", "offset": match.start(), "length": len(match.group()),
             "type": {"@type": "textEntityTypeUrl" if match.group().startswith("http") else "textEntityTypeMention"}}
            for match in _ENTITY.finditer(text)]


@dataclass(frozen=True)
class _Sent:
    """A message the demo account sent through an approved draft."""
    chat_id: int
    message_id: int
    text: str
    document: str | None
    date: datetime


class DemoProvider:
    """The few provider calls the real artifact readers and the target resolver make."""

    def artifact_account(self, expected=None):
        return account.ACCOUNT_ID

    def request_budget(self, deadline):
        return nullcontext()

    def ensure_ready(self):
        return None

    def get_account_id(self):
        return account.ACCOUNT_ID

    def get_self_chat(self):
        return _tdlib_chat(account.CHAT_BY_ID[account.SAVED])

    def resolve_target(self, target):
        name = str(target).removeprefix("@").casefold()
        for chat in account.CHATS:
            if chat.username is not None and chat.username.casefold() == name:
                return _tdlib_chat(chat)
        raise MessageNotFound("no such synthetic username")


def _tdlib_chat(chat: account.Chat) -> dict:
    kinds = {"self": {"@type": "chatTypePrivate", "user_id": account.ACCOUNT_ID},
             "private": {"@type": "chatTypePrivate", "user_id": chat.chat_id},
             "supergroup": {"@type": "chatTypeSupergroup", "supergroup_id": -chat.chat_id, "is_channel": False},
             "channel": {"@type": "chatTypeSupergroup", "supergroup_id": -chat.chat_id, "is_channel": True}}
    return {"@type": "chat", "id": chat.chat_id, "title": chat.title, "type": kinds[chat.kind]}


# --- the fake service ------------------------------------------------------------------------------

@dataclass
class _Draft:
    draft_id: str
    kind: str  # text or document
    recipient: int
    recipient_title: str
    expires_at: datetime
    sha256: str
    size_bytes: int
    text: str | None = None
    artifact_id: str | None = None
    display_name: str | None = None
    mime_type: str | None = None
    caption: str | None = None
    reply_source: object | None = None
    reply_preview: object | None = None  # ReplyDraftPreview or ReplyArtifactDraftPreview
    receipt: object | None = None  # the replayed send response
    sent_message_id: int | None = None
    cancelled: bool = False

    @property
    def summary(self) -> DraftSummary:
        return DraftSummary(draft_id=self.draft_id, account_id=account.ACCOUNT_ID, recipient=self.recipient,
                            recipient_title=self.recipient_title, kind=self.kind, expires_at=self.expires_at)

    @property
    def preview(self) -> DraftPreview:
        content = dict(sha256=self.sha256, size_bytes=self.size_bytes)
        if self.kind == "text":
            content["text"] = self.text
        else:
            content.update(artifact_id=self.artifact_id, display_name=self.display_name, mime_type=self.mime_type,
                           caption=self.caption)
        return DraftPreview(**self.summary.model_dump(), **content)


class DemoService:
    """In-process fake of the broker client over the synthetic account."""

    def __init__(self, *, policy: RuntimePolicy, workdir: Path, transcript: Transcript, approval: str = "approve",
                 send_outcome: str = "sent", status_after_unknown: str = "sent"):
        if approval not in APPROVALS or send_outcome not in {"sent", "outcome_unknown"} or \
                status_after_unknown not in {"sent", "outcome_unknown"}:
            raise ValueError("unknown demo script")
        self.policy = policy
        self.transcript = transcript
        self.approval = approval
        self.send_outcome = send_outcome
        self.status_after_unknown = status_after_unknown
        self.provider = DemoProvider()
        self.store = ArtifactStore(cache_dir=workdir / "cache")
        self.staging = workdir / "staging"
        self.staging.mkdir(mode=0o700, exist_ok=True)
        self.files = {key: synthetic_file(document) for key, document in account.DOCUMENTS.items()}
        self.sent: list[_Sent] = []
        self.artifacts: dict[str, tuple] = {}  # artifact_id -> (anchor or None, name, mime, media kind)
        self.drafts: dict[str, _Draft] = {}
        self.cursors: dict[str, tuple] = {}
        self.lock = threading.RLock()
        self.now = datetime.now(UTC).replace(microsecond=0)
        self.generation = "broker_" + secrets.token_hex(16)
        readers = dict(client=self.provider, store=self.store, metadata_lookup=self.artifacts.get,
                       client_id="client_demo_conversation_harness", broker_generation=self.generation)
        self.pages = AttachmentPageReader(**readers)
        self.sheets = SpreadsheetReader(**readers)
        self.slides = PresentationReader(**readers)
        self.resolver = TargetResolver(self.provider)

    def close(self):
        for reader in (self.pages, self.sheets, self.slides):
            reader.close()

    def __getattr__(self, name):
        if name.startswith("_"):
            raise AttributeError(name)

        def unsupported(*_args, **_kwargs):
            raise ToolError(UNSUPPORTED)
        return unsupported

    # helpers ---------------------------------------------------------------------------------------

    def _messages(self, chat_id: int) -> list:
        """The chat's messages, newest first, including messages the demo sent."""
        rows = [m for m in account.MESSAGES if m.chat_id == chat_id]
        rows += [s for s in self.sent if s.chat_id == chat_id]
        return sorted(rows, key=lambda m: m.message_id, reverse=True)

    def _message(self, chat_id: int, message_id: int):
        return next((m for m in self._messages(chat_id) if m.message_id == message_id), None)

    @staticmethod
    def _sender(m) -> tuple[str, int]:
        if isinstance(m, _Sent):
            return "user", account.ACCOUNT_ID
        return ("chat", m.sender_id) if m.sender_id < 0 else ("user", m.sender_id)

    def _sender_name(self, m) -> str:
        kind, identity = self._sender(m)
        return account.CHAT_BY_ID[identity].title if kind == "chat" else account.USERS[identity]

    @staticmethod
    def _link(m) -> str | None:
        chat = account.CHAT_BY_ID[m.chat_id]
        return safe_telegram_link(f"https://t.me/c/{chat.link_id}/{m.message_id}") if chat.link_id else None

    def _source(self, m) -> SourceEvidence:
        return SourceEvidence(telegram_url=self._link(m), evidence_anchor=EvidenceAnchor(chat_id=m.chat_id,
                                                                                         message_id=m.message_id))

    def _selected(self, m) -> SelectedMessage:
        kind, identity = self._sender(m)
        topic = getattr(m, "topic", None)
        return SelectedMessage(
            content_kind="document" if m.document else "text", text_role="caption" if m.document else "text",
            text=SelectedMessageText(value=m.text, sanitized=False, truncated=False, original_characters=len(m.text)),
            sender=SelectedMessageSender(kind=kind, id=identity, display_name=self._sender_name(m)),
            is_outgoing=bool(getattr(m, "outgoing", True)), date_utc=m.date, edit_date_utc=None, is_edited=False,
            reply_to=None, topic=SelectedMessageTopic(kind="forum", id=topic) if topic else None,
            media_album_id=None, source=self._source(m))

    def _read_result(self, m) -> MessageReadResult:
        return MessageReadResult(anchor=SelectedMessageAnchor(chat_id=m.chat_id, message_id=m.message_id),
                                 status="complete", coverage_complete=True, message=self._selected(m))

    def _date_to(self) -> datetime:
        return max(self.now, account.BASE + timedelta(minutes=1))

    def _cursor(self, prefix: str, state: tuple) -> str:
        token = prefix + secrets.token_hex(32)
        self.cursors[token] = state
        return token

    def _take_cursor(self, token: str | None, facts: tuple):
        if token is None:
            return None
        state = self.cursors.pop(token, None)
        if state is None or state[0] != facts:
            raise LookupError("cursor")
        return state

    def _window(self, rows, request):
        if request.mode == "interval":
            return [m for m in rows if request.date_from <= m.date < request.date_to]
        return rows

    def _expires(self) -> datetime:
        return self.now + timedelta(seconds=300)

    # connection ------------------------------------------------------------------------------------

    def diagnostics(self) -> dict:
        return {"status": "compatible", "broker": {"broker_generation": self.generation, "service": "synthetic demo"},
                "detail": "synthetic demo service; no Telegram connection"}

    # targets ---------------------------------------------------------------------------------------

    def resolve(self, request):
        return self.resolver.resolve(request)

    def discover(self, request):
        lists = ["main", "archive"] if request.scope == "both" else [request.scope]
        facts = (tuple(request.hypotheses), request.scope)
        try:
            state = self._take_cursor(request.cursor, facts)
        except LookupError:
            return TargetDiscoveryResponse(status="expired", coverage=self._discovery_coverage(
                request, lists, lane="expired"))
        page = 1 if state is None else state[1] + 1
        emitted = set() if state is None else state[2]
        candidates = []
        for chat in account.CHATS:
            if chat.chat_list not in lists or chat.chat_id in emitted:
                continue
            if page == 1 and chat.chat_list != lists[0]:
                continue  # the first page scans only the first requested list
            metadata, evidence = [], []
            title = normalize_search_text(chat.title)
            for index, hypothesis in enumerate(request.hypotheses):
                value = normalize_search_text(hypothesis)
                names = {title, normalize_search_text(chat.username or "")} - {""}
                if value in names:
                    metadata.append(MetadataEvidence(hypothesis_index=index, kind="normalized_title"))
                elif any(value in name or name in value for name in names) or \
                        set(value.split()) & set(title.split()) - {"the", "a", "of"}:
                    metadata.append(MetadataEvidence(hypothesis_index=index, kind="title_substring"))
                for m in self._messages(chat.chat_id):
                    if len(evidence) < 3 and len(value) >= 3 and value in normalize_search_text(m.text):
                        evidence.append(GlobalMessageEvidence(
                            hypothesis_index=index, chat_list=chat.chat_list, message_id=m.message_id,
                            date_utc=m.date, snippet=render_evidence(m.text, max_length=512),
                            evidence_anchor=EvidenceAnchor(chat_id=chat.chat_id, message_id=m.message_id)))
            if metadata or evidence:
                candidates.append(DiscoveryCandidate(
                    chat_id=chat.chat_id, title=chat.title, chat_type=chat.kind, list_membership=[chat.chat_list],
                    metadata_evidence=metadata[:10], message_evidence=evidence))
        budget, kept = 10, []
        for candidate in candidates:
            candidate.message_evidence = candidate.message_evidence[:budget]
            budget -= len(candidate.message_evidence)
            kept.append(candidate)
        done = page >= len(lists)
        emitted = emitted | {candidate.chat_id for candidate in kept}
        return TargetDiscoveryResponse(
            status="complete" if done else "page", candidates=kept,
            coverage=self._discovery_coverage(request, lists, lane="complete" if done else "scanning"),
            next_cursor=self._cursor("scan_", (facts, page, emitted)))

    def _discovery_coverage(self, request, lists, *, lane: str) -> DiscoveryCoverage:
        def catalog(name):
            if name not in lists:
                return CatalogLaneCoverage(status="not_requested", scanned_count=0, emitted_count=0, end_reached=False)
            scanned = sum(1 for chat in account.CHATS if chat.chat_list == name) \
                if lane == "complete" or (lane == "scanning" and name == lists[0]) else 0
            return CatalogLaneCoverage(status=lane, scanned_count=scanned, emitted_count=0,
                                       end_reached=lane == "complete")
        global_status = {"complete": "complete", "scanning": "scanning", "expired": "not_started"}[lane]
        lanes = [GlobalMessageLaneCoverage(hypothesis_index=index, chat_list=name, status=global_status,
                                           pages_scanned=0 if lane == "expired" else 1, hits_seen=0)
                 for index in range(len(request.hypotheses)) for name in lists]
        return DiscoveryCoverage(complete=lane == "complete", catalog=CatalogCoverage(main=catalog("main"),
                                 archive=catalog("archive")), global_messages=GlobalMessageCoverage(lanes=lanes),
                                 hydration="complete", detail="discovery coverage details redacted")

    def _chat_for(self, target) -> account.Chat | None:
        if isinstance(target, str):
            name = target.removeprefix("@").casefold()
            return next((c for c in account.CHATS if c.username and c.username.casefold() == name), None)
        return account.CHAT_BY_ID.get(target)

    # search and read -------------------------------------------------------------------------------

    def search(self, request):
        chat = self._chat_for(request.target)
        query = request.query
        if chat is None:
            return SearchResponse(status="error", coverage=Coverage(
                complete=False, text_status="not_started", metadata_status="not_started",
                date_from=request.date_from, date_to=request.date_to, detail="target could not be resolved"),
                matches=[])
        text_requested = query.text is not None or query.contains_number is True
        metadata_requested = any(v is not None for v in (query.file_name, query.mime_type, query.media_type))
        rows = self._messages(chat.chat_id)
        matches = []
        for m in rows:
            if request.date_from and m.date < request.date_from or request.date_to and m.date > request.date_to:
                continue
            document = account.DOCUMENTS.get(m.document) if m.document else None
            body = normalize_search_text(m.text)
            if query.text is not None and not all(word in body for word in normalize_search_text(query.text).split()):
                continue
            if query.contains_number and not re.search(r"\d", m.text):
                continue
            if metadata_requested:
                if document is None:
                    continue
                if query.file_name and normalize_search_text(query.file_name) not in normalize_search_text(
                        document.file_name):
                    continue
                if query.mime_type and query.mime_type.casefold() != document.mime_type.casefold():
                    continue
                if query.media_type and query.media_type != "document":
                    continue
            position = rows.index(m)
            neighbours = rows[max(0, position - request.context_messages):position] + \
                rows[position + 1:position + 1 + request.context_messages] if request.context_messages else []
            matches.append(SearchMatch(
                chat_id=m.chat_id, message_id=m.message_id, date_utc=m.date, sender=self._sender_name(m),
                snippet=render_evidence(m.text),
                context=[ContextEvidence(message_id=n.message_id, date_utc=n.date, sender=self._sender_name(n),
                                         snippet=render_evidence(n.text))
                         for n in sorted(neighbours, key=lambda n: n.message_id)],
                file=FileEvidence(name=document.file_name, mime_type=document.mime_type, media_type="document",
                                  size=len(self.files[document.key])) if document else None,
                source=self._source(m)))
        return SearchResponse(
            status="matches" if matches else "no_match",
            coverage=Coverage(complete=True, text_status="complete" if text_requested else "not_requested",
                              metadata_status="complete" if metadata_requested else "not_requested",
                              date_from=request.date_from, date_to=request.date_to,
                              detail="complete for the exact target and requested date range"),
            matches=matches[:request.limit])

    def get_message_context(self, request):
        anchor = request.anchor
        rows = sorted(self._messages(anchor.chat_id), key=lambda m: m.message_id)
        center = next((i for i, m in enumerate(rows) if m.message_id == anchor.message_id), None)
        if center is None:
            return MessageContextResponse(status="not_found", anchor=anchor, coverage_complete=False,
                                          detail="message is unavailable")
        window = rows[max(0, center - request.radius):center + request.radius + 1]
        return MessageContextResponse(
            status="complete", anchor=anchor, coverage_complete=True, detail="bounded context around exact anchor",
            messages=[ContextEvidence(message_id=m.message_id, date_utc=m.date, sender=self._sender_name(m),
                                      snippet=render_evidence(m.text)) for m in window])

    def read_messages(self, request):
        results = []
        for anchor in request.anchors:
            m = self._message(anchor.chat_id, anchor.message_id)
            results.append(self._read_result(m) if m is not None else MessageReadResult(
                anchor=anchor, status="not_found", coverage_complete=False, issues=["message_unavailable"]))
        complete = all(r.coverage_complete for r in results)
        return ReadMessagesResponse(status="complete" if complete else "partial", coverage_complete=complete,
                                    results=results)

    def read_history(self, request):
        chat = account.CHAT_BY_ID.get(request.target)
        facts = (request.target, request.mode, request.date_from, request.date_to, request.limit)
        try:
            state = self._take_cursor(request.cursor, facts)
        except LookupError:
            return ReadHistoryResponse(status="invalid_cursor", stop_reason="invalid_cursor")
        if chat is None:
            return ReadHistoryResponse(status="error", stop_reason="provider_error")
        rows = self._window(self._messages(chat.chat_id), request) if state is None else state[1]
        scope = state[2] if state is not None else HistoryScope(
            target=chat.chat_id, mode=request.mode, date_from=request.date_from,
            date_to=request.date_to if request.mode == "interval" else self._date_to(),
            upper_message_id=rows[0].message_id if rows else None,
            candidate_limit=100 if request.mode == "latest" else 1000, page_limit=request.limit,
            expires_at=self._expires())
        page, rest = rows[:request.limit], rows[request.limit:]
        results = [self._read_result(m) for m in page]
        cursor = self._cursor("history_", (facts, rest, scope)) if rest else None
        return ReadHistoryResponse(
            status="page" if cursor else "partial", scope=scope, results=results,
            page_complete=len(results) == scope.page_limit, has_more=True if cursor else None, next_cursor=cursor,
            stop_reason="page_limit" if cursor else "provider_end_unverified",
            scanned_candidates=len(results))

    def _topic_result(self, chat: account.Chat, topic_id: int) -> ForumTopicResult | None:
        name = dict(chat.topics).get(topic_id)
        if name is None:
            return None
        created = min((m.date for m in account.MESSAGES if m.chat_id == chat.chat_id and m.topic == topic_id),
                      default=account.BASE)
        return ForumTopicResult(chat_id=chat.chat_id, topic=ForumTopicReference(id=topic_id), status="complete",
                                coverage_complete=True, metadata=ForumTopicMetadata(
                                    name=ForumTopicName(value=name, sanitized=False, truncated=False),
                                    creation_date_utc=created, is_general=topic_id == 1, is_closed=False,
                                    is_hidden=False))

    def list_topics(self, request):
        chat = account.CHAT_BY_ID.get(request.target)
        if request.cursor is not None:
            return ListTopicsResponse(status="invalid_cursor", stop_reason="invalid_cursor")
        if chat is None or not chat.forum:
            return ListTopicsResponse(status="unsupported", stop_reason="unsupported_forum")
        results = [self._topic_result(chat, topic_id) for topic_id, _ in chat.topics][:request.limit]
        return ListTopicsResponse(
            status="partial", scope=TopicScope(target=chat.chat_id, limit=request.limit, expires_at=self._expires()),
            results=results, page_complete=bool(results), stop_reason="provider_end_unverified",
            scanned_candidates=len(results), provider_pages=1)

    def read_topic_history(self, request):
        chat = account.CHAT_BY_ID.get(request.target)
        if request.cursor is not None:
            return ReadTopicHistoryResponse(status="invalid_cursor", stop_reason="invalid_cursor")
        if chat is None or not chat.forum:
            return ReadTopicHistoryResponse(status="unsupported", stop_reason="unsupported_forum")
        topic = self._topic_result(chat, request.topic.id)
        if topic is None:
            return ReadTopicHistoryResponse(status="error", stop_reason="topic_unavailable")
        rows = [m for m in self._window(self._messages(chat.chat_id), request)
                if getattr(m, "topic", None) == request.topic.id][:request.limit]
        scope = TopicHistoryScope(target=chat.chat_id, topic=request.topic, mode=request.mode,
                                  date_from=request.date_from,
                                  date_to=request.date_to if request.mode == "interval" else self._date_to(),
                                  upper_message_id=rows[0].message_id if rows else None, page_limit=request.limit,
                                  expires_at=self._expires())
        results = [self._read_result(m) for m in rows]
        return ReadTopicHistoryResponse(
            status="partial", scope=scope, topic=topic, results=results, page_complete=bool(results),
            stop_reason="provider_end_unverified", scanned_candidates=len(results),
            processed_candidates=len(results), provider_pages=1)

    def search_messages(self, request):
        chat = account.CHAT_BY_ID.get(request.target)
        if request.cursor is not None:
            return SearchMessagesResponse(status="invalid_cursor", stop_reason="invalid_cursor")
        if chat is None:
            return SearchMessagesResponse(status="error", stop_reason="provider_error")
        if request.topic is not None and (not chat.forum or request.topic.id not in dict(chat.topics)):
            return SearchMessagesResponse(status="unsupported" if not chat.forum else "error",
                                          stop_reason="unsupported_forum" if not chat.forum else "topic_unavailable")
        rows = [m for m in self._window(self._messages(chat.chat_id), request)
                if (request.sender is None or self._sender(m) == (request.sender.kind, request.sender.id))
                and (request.direction is None or bool(getattr(m, "outgoing", True)) == (request.direction == "outgoing"))
                and (request.topic is None or getattr(m, "topic", None) == request.topic.id)]

        def searchable(m):
            document = account.DOCUMENTS.get(m.document) if m.document else None
            return m.text + (" " + document.file_name if document else "")

        def lexical(query: str):
            words = normalize_search_text(query).split()
            return [m for m in rows if all(word in normalize_search_text(searchable(m)) for word in words)]

        scope = SearchMessagesScope(
            target=chat.chat_id, query=request.query, sender=request.sender, direction=request.direction,
            topic=request.topic, mode=request.mode, date_from=request.date_from,
            date_to=request.date_to if request.mode == "interval" else self._date_to(),
            upper_message_id=self._messages(chat.chat_id)[0].message_id, page_limit=request.limit,
            expires_at=self._expires())
        branches = []
        if isinstance(request.query, BooleanSearchQuery):
            seen = {}
            for index, seed in enumerate(boolean_seeds(request.query)):
                hits = lexical(seed)[:20]
                branches.append(SearchBranchCoverage(index=index, query=seed, state="provider_end_unverified",
                                                     scanned_candidates=len(hits), processed_candidates=len(hits),
                                                     provider_pages=1))
                for m in hits:
                    seen[m.message_id] = m
            matched = [m for m in sorted(seen.values(), key=lambda m: m.message_id, reverse=True)
                       if boolean_matches(request.query, m.text)]
            scanned = sum(b.scanned_candidates for b in branches)
            pages = len(branches)
        else:
            hits = lexical(request.query)[:20]
            matched, scanned, pages = hits, len(hits), 1
        results = [SearchMessageResult(anchor=SelectedMessageAnchor(chat_id=m.chat_id, message_id=m.message_id),
                                       status="match", coverage_complete=True, message=self._selected(m))
                   for m in matched[:request.limit]]
        return SearchMessagesResponse(
            status="partial", scope=scope, results=results, page_complete=bool(results),
            stop_reason="provider_end_unverified", scanned_candidates=scanned, processed_candidates=scanned,
            provider_pages=pages, branch_coverage=branches)

    def list_chats(self, request):
        lists = ["main", "archive"] if request.scope == "both" else [request.scope]
        facts = (request.scope, request.limit)
        try:
            state = self._take_cursor(request.cursor, facts)
        except LookupError:
            return ListChatsResponse(status="invalid_cursor", stop_reason="invalid_cursor")
        scope = state[2] if state is not None else ChatListScope(
            selection=request.scope, selected_lists=lists, limit=request.limit,
            prefix_limit_per_list=100 if request.scope == "both" else 200, expires_at=self._expires())
        ordered = [(name, rank, chat) for name in lists
                   for rank, chat in enumerate((c for c in account.CHATS if c.chat_list == name), 1)]
        offset = 0 if state is None else state[1]
        page = ordered[offset:offset + request.limit]
        processed = offset + len(page)
        coverage = []
        for name in lists:
            observed = sum(1 for row in ordered if row[0] == name)
            done = sum(1 for row in ordered[:processed] if row[0] == name)
            coverage.append(ChatListCoverage(chat_list=name, native_state="observed", observed=observed,
                                             processed=done, returned=done, pending=observed - done))
        results = [ChatListResult(
            chat_id=chat.chat_id, kind="private" if chat.kind == "self" else chat.kind,
            title=ForumTopicName(value=chat.title, sanitized=False, truncated=False), observed_list=name,
            observed_rank=rank, current_lists=[name], unread_count=account.unread_count(chat.chat_id),
            unread_mention_count=0, unread_reaction_count=0, is_marked_as_unread=False, hydrated_at=self.now)
            for name, rank, chat in page]
        cursor = self._cursor("chats_", (facts, processed, scope)) if processed < len(ordered) else None
        return ListChatsResponse(
            status="page" if cursor else "prefix_exhausted_unverified", scope=scope, results=results,
            list_coverage=coverage, observed_candidates=len(ordered), processed_candidates=processed,
            returned_candidates=processed, pending_candidates=len(ordered) - processed,
            processed_this_page=len(results), returned_this_page=len(results), snapshot_complete=True,
            page_complete=bool(results), next_cursor=cursor,
            stop_reason="page_limit" if cursor else "prefix_exhausted_unverified")

    # attachments -----------------------------------------------------------------------------------

    def _own(self, data: bytes, name: str, mime: str, anchor) -> object:
        staged = Path(tempfile.mkdtemp(dir=self.staging)) / "file"
        try:
            staged.write_bytes(data)
            staged.chmod(0o600)
            artifact = self.store.store(staged, kind="document")
        finally:
            shutil.rmtree(staged.parent, ignore_errors=True)
        with self.lock:
            self.artifacts[artifact.artifact_id] = (anchor, name, mime, "document")
        return artifact

    def get_attachment(self, request):
        anchor = request.anchor
        m = self._message(anchor.chat_id, anchor.message_id)
        if m is None:
            return AttachmentResponse(status="not_found", anchor=anchor, detail="message is unavailable")
        if not m.document:
            return AttachmentResponse(status="unsupported", anchor=anchor, detail="message has no supported attachment")
        document = account.DOCUMENTS[m.document]
        artifact = self._own(self.files[document.key], document.file_name, document.mime_type, anchor)
        return AttachmentResponse(
            status="complete", anchor=anchor, media_type="document", file_name=document.file_name,
            mime_type=document.mime_type, size_bytes=artifact.size_bytes, sha256=artifact.sha256,
            artifact_id=artifact.artifact_id, artifact_path=str(artifact.path),
            expires_at=datetime.fromtimestamp(artifact.expires_at, UTC), coverage="complete",
            detail="original bytes are available in the private cache")

    def create_local_artifact(self, request):
        import base64
        data = request.content.encode("utf-8") if request.content is not None else \
            base64.b64decode(request.content_base64, validate=True)
        mime = "text/plain" if request.file_name.endswith(".txt") else "application/octet-stream"
        artifact = self._own(data, request.file_name, mime, None)
        return CreateLocalArtifactResponse(status="complete", artifact_id=artifact.artifact_id,
                                           artifact_path=str(artifact.path), sha256=artifact.sha256,
                                           size_bytes=artifact.size_bytes, file_name=request.file_name,
                                           detail="local artifact created; nothing was sent")

    def read_attachment(self, request):
        metadata = self.artifacts.get(request.artifact_id)
        artifact = self.store.lookup(request.artifact_id)
        if metadata is None or artifact is None:
            return ReadAttachmentResponse(status="expired", detail="artifact metadata is unavailable")
        anchor, name, mime, _ = metadata
        result = read_document(artifact.path, name=name, mime_type=mime, max_chars=request.max_chars,
                               max_pages=request.max_pages)
        images = []
        for index, path in enumerate(result.image_paths):
            stored = self.store.store(path, kind="image") if path != artifact.path else artifact
            images.append(AttachmentImage(artifact_id=stored.artifact_id, artifact_path=str(stored.path),
                                          mime_type="image/png",
                                          page_number=index + 1 if result.total_pages is not None else None))
        for path in result.image_paths:
            if path != artifact.path:
                path.unlink(missing_ok=True)
                shutil.rmtree(path.parent, ignore_errors=True)
        return ReadAttachmentResponse(
            status=result.status, source_anchor=anchor, text=result.text, processed_bytes=result.processed_bytes,
            total_bytes=result.total_bytes, processed_pages=result.processed_pages, total_pages=result.total_pages,
            images=images[:5], coverage_complete=result.status == "complete",
            detail=result.error or "bounded local document read")

    def read_attachment_page(self, request):
        return self.pages.read(request)

    def read_spreadsheet(self, request):
        return self.sheets.read(request)

    def read_presentation(self, request):
        return self.slides.read(request)

    # drafts ----------------------------------------------------------------------------------------

    def _new_draft(self, **fields) -> _Draft:
        draft = _Draft(draft_id="draft_" + secrets.token_hex(16), expires_at=self.now + DRAFT_TTL, **fields)
        with self.lock:
            self.drafts[draft.draft_id] = draft
        return draft

    def _pending(self, draft_id: str, kind: str | None = None, reply: bool | None = None) -> _Draft | None:
        draft = self.drafts.get(draft_id)
        if draft is None or draft.cancelled or draft.receipt is not None or \
                (kind is not None and draft.kind != kind) or (reply is not None and (draft.reply_source is not None) != reply):
            return None
        return draft

    def prepare_text_send(self, request):
        chat = account.CHAT_BY_ID.get(request.recipient)
        if request.topic is not None or request.outgoing_entities:
            return PrepareTextSendResponse(status="error", detail="topic draft could not be prepared safely")
        if chat is None:
            return PrepareTextSendResponse(status="error", detail="text draft could not be prepared safely")
        text = _safe_message_text(request.text)
        data = text.encode("utf-8")
        draft = self._new_draft(kind="text", recipient=chat.chat_id, recipient_title=chat.title, text=text,
                                sha256=hashlib.sha256(data).hexdigest(), size_bytes=len(data), display_name="Text message",
                                mime_type="text/plain", caption="")
        return PrepareTextSendResponse(
            status="prepared", draft_id=draft.draft_id, recipient=chat.chat_id, recipient_title=chat.title, text=text,
            sha256=draft.sha256, expires_at=draft.expires_at,
            detail="immutable text draft prepared; explicit approval is required before sending")

    def _artifact_fields(self, request) -> dict | None:
        if request.artifact_id not in self.artifacts or request.kind != "document" or \
                self.store.lookup(request.artifact_id) is None or not _MIME_TYPE.fullmatch(request.mime_type):
            return None
        parts = request.artifact_id.split("_")
        return dict(artifact_id=request.artifact_id, display_name=_safe_display_name(request.display_name),
                    mime_type=request.mime_type, caption=_safe_caption(request.caption), sha256=parts[2],
                    size_bytes=int(parts[3]))

    def prepare_artifact_send(self, request):
        chat = account.CHAT_BY_ID.get(request.recipient)
        fields = self._artifact_fields(request)
        if chat is None or fields is None:
            return PrepareArtifactSendResponse(status="error", detail="draft could not be prepared safely")
        draft = self._new_draft(kind="document", recipient=chat.chat_id, recipient_title=chat.title, **fields)
        return PrepareArtifactSendResponse(
            status="prepared", draft_id=draft.draft_id, recipient=chat.chat_id, recipient_title=chat.title,
            display_name=draft.display_name, mime_type=draft.mime_type, caption=draft.caption, kind="document",
            sha256=draft.sha256, size_bytes=draft.size_bytes, expires_at=draft.expires_at,
            detail="immutable local draft prepared; explicit approval is required before sending")

    def _reply_source(self, request, tool_capability: str, options: tuple[str, ...]):
        """(source, None) or (None, refusal kwargs): the real projection names the source's families."""
        anchor = request.reply_to
        m = self._message(anchor.chat_id, anchor.message_id) if anchor is not None else None
        if m is None or anchor.chat_id != request.recipient or isinstance(m, _Sent):
            return None, dict(status="unavailable", detail=UNAVAILABLE)
        topic = getattr(request, "topic", None)
        try:
            source = source_from_message(self._tdlib_message(m), m.chat_id, m.message_id,
                                         expected_forum_topic_id=topic.id if topic is not None else None)
        except ValueError:
            return None, dict(status="unavailable", detail=UNAVAILABLE)
        required = (tool_capability, "send", *options, *source.required_capabilities)
        missing = [name for name in dict.fromkeys(required) if name not in self.policy.enabled_capabilities]
        if missing:
            return None, dict(status="capability_disabled", detail=CAPABILITY_DISABLED, missing_capabilities=missing)
        return source, None

    def _tdlib_message(self, m) -> dict:
        kind, identity = self._sender(m)
        caption = {"@type": "formattedText", "text": m.text, "entities": _entities(m.text)}
        if m.document:
            document = account.DOCUMENTS[m.document]
            size = len(self.files[document.key])
            unique = "demo-" + document.key
            content = {"@type": "messageDocument", "caption": caption, "document": {
                "@type": "document", "file_name": document.file_name, "mime_type": document.mime_type,
                "document": {"@type": "file", "id": m.message_id, "size": size, "expected_size": size,
                             "local": {"@type": "localFile", "path": "", "can_be_downloaded": True,
                                       "can_be_deleted": False, "is_downloading_active": False,
                                       "is_downloading_completed": False, "download_offset": 0,
                                       "downloaded_prefix_size": 0, "downloaded_size": 0},
                             "remote": {"@type": "remoteFile", "id": unique, "unique_id": unique,
                                        "is_uploading_active": False, "is_uploading_completed": True,
                                        "uploaded_size": size}}}}
        else:
            content = {"@type": "messageText", "text": caption}
        sender = {"@type": "messageSenderUser", "user_id": identity} if kind == "user" else \
            {"@type": "messageSenderChat", "chat_id": identity}
        topic = getattr(m, "topic", None)
        return {"@type": "message", "chat_id": m.chat_id, "id": m.message_id, "sender_id": sender,
                "is_outgoing": bool(getattr(m, "outgoing", True)), "is_from_offline": False,
                "ephemeral_message_id": 0, "date": int(m.date.timestamp()), "edit_date": 0, "self_destruct_in": 0,
                "auto_delete_in": 0, "content": content,
                "topic_id": {"@type": "messageTopicForum", "forum_topic_id": topic} if topic else None}

    def prepare_reply_text_send(self, request):
        options = (("formatted_text_send",) if request.outgoing_entities else ()) + \
            (("topic_text_send",) if request.topic is not None else ())
        source, refusal = self._reply_source(request, "reply_text_send", options)
        if refusal is not None:
            return PrepareReplyTextSendResponse(**refusal)
        if options:  # styles and topic sends are outside the demo policy, so the guard above refused them
            return PrepareReplyTextSendResponse(status="unavailable", detail=UNAVAILABLE)
        chat = account.CHAT_BY_ID[request.recipient]
        text = _safe_message_text(request.text)
        data = text.encode("utf-8")
        draft = self._new_draft(kind="text", recipient=chat.chat_id, recipient_title=chat.title, text=text,
                                sha256=hashlib.sha256(data).hexdigest(), size_bytes=len(data),
                                display_name="Text message", mime_type="text/plain", caption="", reply_source=source)
        inner = ReplyTextDraftPreview(draft_id=draft.draft_id, account_id=account.ACCOUNT_ID, recipient=chat.chat_id,
                                      recipient_title=chat.title, kind="text", expires_at=draft.expires_at,
                                      sha256=draft.sha256, size_bytes=draft.size_bytes, text=text)
        draft.reply_preview = ReplyDraftPreview.create(inner, source)
        return PrepareReplyTextSendResponse(status="prepared", reply=draft.reply_preview, detail=PREPARED)

    def prepare_reply_artifact_send(self, request):
        source, refusal = self._reply_source(request, "reply_artifact_send", ())
        if refusal is not None:
            return PrepareReplyArtifactSendResponse(**refusal)
        fields = self._artifact_fields(request)
        if fields is None:
            return PrepareReplyArtifactSendResponse(status="unavailable", detail=UNAVAILABLE)
        chat = account.CHAT_BY_ID[request.recipient]
        draft = self._new_draft(kind="document", recipient=chat.chat_id, recipient_title=chat.title,
                                reply_source=source, **fields)
        draft.reply_preview = ReplyArtifactDraftPreview.create(draft.preview, source)
        return PrepareReplyArtifactSendResponse(status="prepared", reply=draft.reply_preview, detail=PREPARED)

    def get_reply_draft(self, request):
        draft = self._pending(request.draft_id, "text", reply=True)
        if draft is None:
            return GetReplyDraftResponse(status="unavailable", detail=UNAVAILABLE)
        return GetReplyDraftResponse(status="pending", reply=draft.reply_preview, detail=PENDING)

    def get_reply_artifact_draft(self, request):
        draft = self._pending(request.draft_id, "document", reply=True)
        if draft is None:
            return GetReplyArtifactDraftResponse(status="unavailable", detail=UNAVAILABLE)
        return GetReplyArtifactDraftResponse(status="pending", reply=draft.reply_preview,
                                             detail="exact pending reply snapshot; inspection is not approval")

    def update_reply_draft(self, request):
        return ReviseReplyDraftResponse(status="unavailable", detail=UNAVAILABLE)

    refresh_reply_draft = update_reply_draft

    def update_reply_artifact_draft(self, request):
        return ReviseReplyArtifactDraftResponse(status="unavailable", detail=UNAVAILABLE)

    refresh_reply_artifact_draft = update_reply_artifact_draft

    def update_draft(self, request):
        return ReviseDraftResponse(status="unavailable", detail="draft is unavailable")

    refresh_draft = update_draft

    def list_drafts(self, request):
        drafts = sorted((d for d in self.drafts.values() if self._pending(d.draft_id) is not None),
                        key=lambda d: d.draft_id)
        drafts = [d for d in drafts if request.after_draft_id is None or d.draft_id > request.after_draft_id]
        page = drafts[:request.limit]
        more = len(drafts) > request.limit
        return ListDraftsResponse(status="listed", drafts=[d.summary for d in page], has_more=more,
                                  next_after_draft_id=page[-1].draft_id if more else None,
                                  detail="owned pending drafts; live view, no recipient verification or approval")

    def get_draft(self, request):
        draft = self._pending(request.draft_id, reply=False)
        if draft is None:
            return GetDraftResponse(status="unavailable", detail="draft is unavailable")
        return GetDraftResponse(status="pending", draft=draft.preview,
                                detail="exact pending draft; no recipient verification or approval")

    def cancel_draft(self, request):
        draft = self._pending(request.draft_id)
        if draft is None:
            return CancelDraftResponse(status="unavailable", detail="draft is unavailable")
        draft.cancelled = True
        self.transcript.write({"event": "cancel", "draft_id": draft.draft_id})
        return CancelDraftResponse(status="cancelled", draft_id=draft.draft_id,
                                   detail="pending draft cancelled; cached artifacts retained")

    # sends -----------------------------------------------------------------------------------------

    def _decide(self, draft: _Draft) -> str:
        """Render the real dialog document for the bound preview, then apply the scripted decision."""
        arguments = dict(recipient_title=draft.recipient_title, recipient=draft.recipient,
                         display_name=draft.display_name or "", size_bytes=draft.size_bytes, sha256=draft.sha256,
                         caption=draft.caption or "", kind=draft.kind, text=draft.text or "")
        if draft.reply_preview is not None:
            key = "reply_preview" if draft.kind == "text" else "reply_artifact_preview"
            arguments[key] = draft.reply_preview.model_dump(mode="json")
        else:
            arguments["draft_preview"] = draft.preview.model_dump(mode="json")
        rendered = render_approval(**arguments)
        decision = APPROVALS[self.approval] if rendered is not None else "dialog_unavailable"
        self.transcript.write({"event": "approval", "draft_id": draft.draft_id, "rendered": rendered is not None,
                               "document_lines": len(rendered[0].splitlines()) if rendered else 0,
                               "decision": decision})
        return decision

    def _send(self, request, *, text_only: bool):
        response = SendPreparedTextResponse if text_only else SendPreparedArtifactResponse
        with self.lock:
            if request.approved is not True:
                return response(status="declined", draft_id=request.draft_id, detail=NOT_GIVEN)
            draft = self.drafts.get(request.draft_id)
            if draft is not None and draft.receipt is not None:
                self.transcript.write({"event": "send_replayed", "draft_id": draft.draft_id})
                return draft.receipt  # a repeated call replays the receipt; it never sends twice
            if draft is None or draft.cancelled:
                return response(status="expired", draft_id=request.draft_id, detail="draft is unavailable or expired")
            if (draft.kind == "text") != text_only:
                return response(status="error", draft_id=request.draft_id, detail="draft kind does not match send tool")
            if draft.reply_source is not None:
                missing = [name for name in draft.reply_source.required_capabilities
                           if name not in self.policy.enabled_capabilities]
                if missing:
                    return response(status="expired", draft_id=request.draft_id, detail="draft is unavailable or expired")
            decision = self._decide(draft)
            if decision != "approved":
                return response(status=decision, draft_id=draft.draft_id, recipient=draft.recipient,
                                detail=APPROVAL_DETAILS[decision])
            message_id = max(m.message_id for m in self._messages(draft.recipient)) + 1
            delivered = self.send_outcome == "sent" or self.status_after_unknown == "sent"
            if delivered:
                document = next((key for key, value in account.DOCUMENTS.items()
                                 if draft.artifact_id and value.file_name == draft.display_name), None)
                self.sent.append(_Sent(draft.recipient, message_id, draft.text or draft.caption or draft.display_name,
                                       document, self._date_to() - timedelta(seconds=1)))
                draft.sent_message_id = message_id
            if self.send_outcome == "sent":
                draft.receipt = response(status="sent", draft_id=draft.draft_id, recipient=draft.recipient,
                                         message_id=message_id, detail="TDLib confirmed delivery")
            else:
                draft.receipt = response(status="outcome_unknown", draft_id=draft.draft_id, recipient=draft.recipient,
                                         detail="delivery is unconfirmed; do not retry automatically")
            self.transcript.write({"event": "send_attempt", "draft_id": draft.draft_id, "recipient": draft.recipient,
                                   "reply_to": list(draft.reply_source.anchor) if draft.reply_source else None,
                                   "outcome": draft.receipt.status, "delivered": delivered})
            return draft.receipt

    def send_prepared_text(self, request):
        return self._send(request, text_only=True)

    def send_prepared_artifact(self, request):
        return self._send(request, text_only=False)

    def get_send_status(self, request):
        draft = self.drafts.get(request.draft_id)
        if draft is None:
            return send_status(request.draft_id, "none")
        if draft.receipt is None:
            return send_status(draft.draft_id, "local_pending")
        if draft.sent_message_id is not None:
            return send_status(draft.draft_id, "provider_confirmed", draft.sent_message_id)
        return send_status(draft.draft_id, "provider_pending", status_override="outcome_unknown")


def serve(arguments: argparse.Namespace, workdir: Path) -> None:
    transcript = Transcript(Path(arguments.transcript) if arguments.transcript else None)
    policy = demo_policy(arguments.policy)
    service = DemoService(policy=policy, workdir=workdir, transcript=transcript, approval=arguments.approval,
                          send_outcome=arguments.send_outcome, status_after_unknown=arguments.status_after_unknown)
    transcript.write({"event": "start", "policy": arguments.policy, "approval": arguments.approval,
                      "send_outcome": arguments.send_outcome, "status_after_unknown": arguments.status_after_unknown,
                      "tools": available_tools(policy.enabled_capabilities)})
    server = build_server(service_factory=lambda: service, artifact_root=workdir / "cache", policy=policy)
    record_tool_calls(server, transcript)
    try:
        server.run("stdio")
    finally:
        transcript.write({"event": "end"})
        service.close()


def parse_arguments(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(prog="conversation_demo_server", description=__doc__.splitlines()[0])
    parser.add_argument("--transcript", help="absolute path of the owner-only JSONL transcript")
    parser.add_argument("--policy", choices=sorted(POLICIES), default="demo")
    parser.add_argument("--approval", choices=sorted(APPROVALS), default="approve")
    parser.add_argument("--send-outcome", choices=("sent", "outcome_unknown"), default="sent")
    parser.add_argument("--status-after-unknown", choices=("sent", "outcome_unknown"), default="sent")
    arguments = parser.parse_args(argv)
    if arguments.transcript and not Path(arguments.transcript).is_absolute():
        parser.error("--transcript must be absolute")
    return arguments


def main(argv=None) -> int:
    arguments = parse_arguments(argv)
    workdir = _WORKDIR or _private_workdir()
    atexit.register(shutil.rmtree, workdir, True)
    serve(arguments, workdir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
